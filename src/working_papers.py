"""
Working Papers Generator Module
Generates audit working papers based on accounting data and documents
"""

import pandas as pd
import logging
from pathlib import Path
from typing import Dict, List, Optional, Any
from datetime import datetime
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows


class WorkingPapersGenerator:
    """Generates audit working papers from accounting data"""
    
    def __init__(self, config: Dict, accounting_processor):
        """
        Initialize working papers generator
        
        Args:
            config: Configuration dictionary
            accounting_processor: AccountingCoreProcessor instance
        """
        self.config = config
        self.accounting = accounting_processor
        self.logger = logging.getLogger(__name__)
        
        # Output configuration
        self.output_dir = Path(config.get('working_papers', {}).get('output_dir', './papeles_trabajo'))
        self.output_dir.mkdir(exist_ok=True)
        
        # Style configuration
        self.header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        self.header_font = Font(name='Arial', size=11, bold=True, color="FFFFFF")
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def generate_balance_analysis(self, year: int, comparison_years: Optional[List[int]] = None) -> Path:
        """
        Generate balance analysis working paper
        
        Args:
            year: Main year to analyze
            comparison_years: Years to compare with
            
        Returns:
            Path to generated Excel file
        """
        self.logger.info(f"Generating balance analysis for {year}")
        
        # Get period
        period = self.accounting.session.query(
            self.accounting.AccountingPeriod
        ).filter_by(year=year).first()
        
        if not period:
            self.logger.error(f"Period {year} not found")
            return None
        
        # Create workbook
        wb = Workbook()
        
        # Sheet 1: Balance at 4 digits minimum
        self._create_balance_sheet(wb, period, level=4)
        
        # Sheet 2: Detailed balance (all levels)
        self._create_balance_sheet(wb, period, level=None, sheet_name="Balance Detallado")
        
        # Sheet 3: Comparative analysis
        if comparison_years:
            self._create_comparative_analysis(wb, year, comparison_years)
        
        # Sheet 4: Account groups summary
        self._create_account_groups_summary(wb, period)
        
        # Sheet 5: Ratios and indicators
        self._create_financial_ratios(wb, period)
        
        # Remove default sheet
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])
        
        # Save file
        filename = f"PT_Balance_{year}_{datetime.now().strftime('%Y%m%d')}.xlsx"
        output_path = self.output_dir / filename
        wb.save(output_path)
        
        self.logger.info(f"Balance analysis saved to {output_path}")
        return output_path
    
    def _create_balance_sheet(self, wb: Workbook, period, level: Optional[int] = None, sheet_name: str = "Balance"):
        """Create balance sheet in workbook"""
        ws = wb.create_sheet(sheet_name)
        
        # Query accounts
        query = self.accounting.session.query(self.accounting.BalanceAccount).filter_by(
            period_id=period.id
        )
        
        if level:
            query = query.filter(self.accounting.BalanceAccount.account_level >= level)
        
        accounts = query.order_by(self.accounting.BalanceAccount.account_code).all()
        
        # Create DataFrame
        data = []
        for acc in accounts:
            data.append({
                'Código': acc.account_code,
                'Nombre': acc.account_name,
                'Suma Debe': acc.suma_debe,
                'Suma Haber': acc.suma_haber,
                'Saldo Deudor': acc.saldo_deudor,
                'Saldo Acreedor': acc.saldo_acreedor,
                'Saldo Final': acc.saldo_final
            })
        
        df = pd.DataFrame(data)
        
        # Write to sheet
        self._write_dataframe_to_sheet(ws, df, f"Balance de Sumas y Saldos - Ejercicio {period.year}")
        
        # Add totals
        if not df.empty:
            last_row = ws.max_row + 1
            ws.cell(last_row, 1, "TOTALES")
            ws.cell(last_row, 1).font = Font(bold=True)
            
            for col_idx, col_name in enumerate(['Suma Debe', 'Suma Haber', 'Saldo Deudor', 'Saldo Acreedor'], start=3):
                total = df[col_name].sum()
                ws.cell(last_row, col_idx, total)
                ws.cell(last_row, col_idx).font = Font(bold=True)
                ws.cell(last_row, col_idx).number_format = '#,##0.00'
    
    def _create_comparative_analysis(self, wb: Workbook, current_year: int, comparison_years: List[int]):
        """Create comparative analysis sheet"""
        ws = wb.create_sheet("Análisis Comparativo")
        
        # Get all periods
        years = [current_year] + comparison_years
        periods = {}
        
        for year in years:
            period = self.accounting.session.query(
                self.accounting.AccountingPeriod
            ).filter_by(year=year).first()
            if period:
                periods[year] = period
        
        if not periods:
            return
        
        # Get accounts at 4 digit level
        all_accounts = set()
        data_by_year = {}
        
        for year, period in periods.items():
            accounts = self.accounting.session.query(
                self.accounting.BalanceAccount
            ).filter_by(period_id=period.id).filter(
                self.accounting.BalanceAccount.account_level >= 4
            ).all()
            
            data_by_year[year] = {acc.account_code: acc for acc in accounts}
            all_accounts.update(acc.account_code for acc in accounts)
        
        # Create comparison table
        headers = ['Código', 'Nombre'] + [str(y) for y in sorted(years)] + ['Variación %']
        ws.append(headers)
        
        # Style header
        for col_idx, _ in enumerate(headers, start=1):
            cell = ws.cell(1, col_idx)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal='center')
        
        # Fill data
        for account_code in sorted(all_accounts):
            row_data = [account_code]
            
            # Get account name from most recent year
            account_name = ""
            for year in sorted(years, reverse=True):
                if account_code in data_by_year[year]:
                    account_name = data_by_year[year][account_code].account_name
                    break
            row_data.append(account_name)
            
            # Add balances for each year
            balances = []
            for year in sorted(years):
                if account_code in data_by_year[year]:
                    balance = data_by_year[year][account_code].saldo_final
                    balances.append(balance)
                    row_data.append(balance)
                else:
                    balances.append(0)
                    row_data.append(0)
            
            # Calculate variation percentage
            if len(balances) >= 2 and balances[-2] != 0:
                variation = ((balances[-1] - balances[-2]) / abs(balances[-2])) * 100
                row_data.append(variation)
            else:
                row_data.append(None)
            
            ws.append(row_data)
        
        # Format columns
        for col in ws.iter_cols(min_col=3, max_col=len(headers)):
            for cell in col:
                if cell.row > 1 and cell.value is not None:
                    cell.number_format = '#,##0.00'
    
    def _create_account_groups_summary(self, wb: Workbook, period):
        """Create account groups summary (Grupo 1-7)"""
        ws = wb.create_sheet("Resumen por Grupos")
        
        # Define account groups
        groups = {
            '1': 'Financiación Básica',
            '2': 'Activo No Corriente',
            '3': 'Existencias',
            '4': 'Acreedores y Deudores',
            '5': 'Cuentas Financieras',
            '6': 'Compras y Gastos',
            '7': 'Ventas e Ingresos'
        }
        
        # Headers
        headers = ['Grupo', 'Descripción', 'Suma Debe', 'Suma Haber', 'Saldo Final']
        ws.append(headers)
        
        # Style header
        for col_idx, _ in enumerate(headers, start=1):
            cell = ws.cell(1, col_idx)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal='center')
        
        # Calculate totals by group
        for group_num, group_name in sorted(groups.items()):
            accounts = self.accounting.session.query(
                self.accounting.BalanceAccount
            ).filter_by(period_id=period.id).filter(
                self.accounting.BalanceAccount.account_code.like(f'{group_num}%')
            ).all()
            
            suma_debe = sum(acc.suma_debe for acc in accounts)
            suma_haber = sum(acc.suma_haber for acc in accounts)
            saldo_final = sum(acc.saldo_final for acc in accounts)
            
            ws.append([group_num, group_name, suma_debe, suma_haber, saldo_final])
        
        # Format numbers
        for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
            for cell in row:
                cell.number_format = '#,##0.00'
    
    def _create_financial_ratios(self, wb: Workbook, period):
        """Create financial ratios sheet"""
        ws = wb.create_sheet("Ratios Financieros")
        
        # Get key accounts
        accounts = self.accounting.session.query(
            self.accounting.BalanceAccount
        ).filter_by(period_id=period.id).all()
        
        # Calculate key balances
        activo_corriente = sum(acc.saldo_final for acc in accounts if acc.account_code.startswith(('3', '4', '5')) and acc.saldo_final > 0)
        activo_no_corriente = sum(acc.saldo_final for acc in accounts if acc.account_code.startswith('2') and acc.saldo_final > 0)
        pasivo_corriente = sum(abs(acc.saldo_final) for acc in accounts if acc.account_code.startswith(('4', '5')) and acc.saldo_final < 0)
        pasivo_no_corriente = sum(abs(acc.saldo_final) for acc in accounts if acc.account_code.startswith('1') and acc.saldo_final < 0 and not acc.account_code.startswith('10'))
        patrimonio_neto = sum(abs(acc.saldo_final) for acc in accounts if acc.account_code.startswith('10'))
        
        # Headers
        ws.append(['Ratio', 'Valor', 'Interpretación'])
        ws.cell(1, 1).font = self.header_font
        ws.cell(1, 2).font = self.header_font
        ws.cell(1, 3).font = self.header_font
        
        # Calculate ratios
        ratios = []
        
        # Liquidez
        if pasivo_corriente > 0:
            ratio_liquidez = activo_corriente / pasivo_corriente
            ratios.append([
                'Ratio de Liquidez',
                ratio_liquidez,
                'Óptimo > 1.5'
            ])
        
        # Solvencia
        total_activo = activo_corriente + activo_no_corriente
        total_pasivo = pasivo_corriente + pasivo_no_corriente
        if total_pasivo > 0:
            ratio_solvencia = total_activo / total_pasivo
            ratios.append([
                'Ratio de Solvencia',
                ratio_solvencia,
                'Óptimo > 1.5'
            ])
        
        # Endeudamiento
        if total_activo > 0:
            ratio_endeudamiento = total_pasivo / total_activo
            ratios.append([
                'Ratio de Endeudamiento',
                ratio_endeudamiento,
                'Óptimo < 0.6'
            ])
        
        # Autonomía financiera
        if patrimonio_neto > 0 and total_pasivo > 0:
            ratio_autonomia = patrimonio_neto / total_pasivo
            ratios.append([
                'Ratio de Autonomía Financiera',
                ratio_autonomia,
                'Óptimo > 0.7'
            ])
        
        # Write ratios
        for ratio_data in ratios:
            ws.append(ratio_data)
        
        # Format
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
            for cell in row:
                cell.number_format = '0.00'
    
    def generate_account_detail(self, account_code: str, year: int) -> Path:
        """
        Generate detailed working paper for specific account
        
        Args:
            account_code: Account code
            year: Year
            
        Returns:
            Path to generated file
        """
        self.logger.info(f"Generating detail for account {account_code} - {year}")
        
        # Get account balance
        period = self.accounting.session.query(
            self.accounting.AccountingPeriod
        ).filter_by(year=year).first()
        
        if not period:
            return None
        
        account = self.accounting.session.query(
            self.accounting.BalanceAccount
        ).filter_by(
            period_id=period.id,
            account_code=account_code
        ).first()
        
        if not account:
            return None
        
        # Get movements
        movements_df = self.accounting.get_account_movements(account_code, year)
        
        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = f"Cuenta {account_code}"
        
        # Account header
        ws['A1'] = 'PAPEL DE TRABAJO - DETALLE DE CUENTA'
        ws['A1'].font = Font(size=14, bold=True)
        
        ws['A3'] = 'Cuenta:'
        ws['B3'] = account_code
        ws['A4'] = 'Nombre:'
        ws['B4'] = account.account_name
        ws['A5'] = 'Ejercicio:'
        ws['B5'] = year
        
        # Balance summary
        ws['A7'] = 'RESUMEN'
        ws['A7'].font = Font(bold=True)
        
        summary_data = [
            ['Saldo Inicial', account.saldo_inicial],
            ['Suma Debe', account.suma_debe],
            ['Suma Haber', account.suma_haber],
            ['Saldo Deudor', account.saldo_deudor],
            ['Saldo Acreedor', account.saldo_acreedor],
            ['Saldo Final', account.saldo_final]
        ]
        
        for i, (label, value) in enumerate(summary_data, start=8):
            ws.cell(i, 1, label)
            ws.cell(i, 2, value)
            ws.cell(i, 2).number_format = '#,##0.00'
        
        # Movements detail
        if not movements_df.empty:
            ws['A15'] = 'MOVIMIENTOS DETALLADOS'
            ws['A15'].font = Font(bold=True)
            
            # Write movements
            start_row = 17
            self._write_dataframe_to_sheet(
                ws,
                movements_df,
                title=None,
                start_row=start_row
            )
        
        # Save
        filename = f"PT_Cuenta_{account_code}_{year}_{datetime.now().strftime('%Y%m%d')}.xlsx"
        output_path = self.output_dir / filename
        wb.save(output_path)
        
        return output_path
    
    def generate_audit_areas_summary(self, year: int) -> Path:
        """
        Generate summary of main audit areas
        
        Args:
            year: Year to analyze
            
        Returns:
            Path to generated file
        """
        wb = Workbook()
        
        # Define audit areas
        audit_areas = {
            'Inmovilizado': ['20', '21', '22', '23'],
            'Existencias': ['30', '31', '32', '33', '34', '35', '36'],
            'Deudores': ['430', '431', '432', '433', '440', '441'],
            'Tesorería': ['57'],
            'Patrimonio Neto': ['10'],
            'Acreedores': ['400', '401', '410', '411'],
            'Ingresos': ['70'],
            'Gastos': ['60', '62', '64']
        }
        
        period = self.accounting.session.query(
            self.accounting.AccountingPeriod
        ).filter_by(year=year).first()
        
        if not period:
            return None
        
        # Create sheet for each area
        for area_name, account_prefixes in audit_areas.items():
            ws = wb.create_sheet(area_name)
            
            # Get accounts for this area
            area_accounts = []
            for prefix in account_prefixes:
                accounts = self.accounting.session.query(
                    self.accounting.BalanceAccount
                ).filter_by(period_id=period.id).filter(
                    self.accounting.BalanceAccount.account_code.like(f'{prefix}%')
                ).all()
                area_accounts.extend(accounts)
            
            if not area_accounts:
                continue
            
            # Create data
            data = []
            for acc in sorted(area_accounts, key=lambda x: x.account_code):
                data.append({
                    'Código': acc.account_code,
                    'Nombre': acc.account_name,
                    'Saldo': acc.saldo_final
                })
            
            df = pd.DataFrame(data)
            self._write_dataframe_to_sheet(ws, df, f"Área: {area_name}")
            
            # Add total
            if not df.empty:
                last_row = ws.max_row + 1
                ws.cell(last_row, 1, "TOTAL")
                ws.cell(last_row, 1).font = Font(bold=True)
                total = df['Saldo'].sum()
                ws.cell(last_row, 3, total)
                ws.cell(last_row, 3).font = Font(bold=True)
                ws.cell(last_row, 3).number_format = '#,##0.00'
        
        # Remove default sheet
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])
        
        # Save
        filename = f"PT_Areas_Auditoria_{year}_{datetime.now().strftime('%Y%m%d')}.xlsx"
        output_path = self.output_dir / filename
        wb.save(output_path)
        
        return output_path
    
    def _write_dataframe_to_sheet(self, ws, df: pd.DataFrame, title: Optional[str] = None, start_row: int = 1):
        """Write DataFrame to Excel sheet with formatting"""
        current_row = start_row
        
        # Write title if provided
        if title:
            ws.cell(current_row, 1, title)
            ws.cell(current_row, 1).font = Font(size=12, bold=True)
            current_row += 2
        
        # Write headers
        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(current_row, col_idx, col_name)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = self.border
        
        current_row += 1
        
        # Write data
        for _, row in df.iterrows():
            for col_idx, value in enumerate(row, start=1):
                cell = ws.cell(current_row, col_idx, value)
                cell.border = self.border
                
                # Format numbers
                if isinstance(value, (int, float)):
                    cell.number_format = '#,##0.00'
        
            current_row += 1
        
        # Adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width
